"""
set_timings.py
Sets 7.5-second (7500ms) advance timing on every slide in NutriAI_Presentation.pptx
so PowerPoint exports exactly 150 seconds of video.
Requires: pip install python-pptx
"""
from pptx import Presentation
from pptx.util import Pt
from pptx.oxml.ns import qn
from lxml import etree
import copy, sys, os

PPTX_PATH = "NutriAI_Presentation.pptx"
SLIDE_DURATION_MS = 7500          # 7.5 s × 20 slides = 150 s

prs = Presentation(PPTX_PATH)

for i, slide in enumerate(prs.slides):
    # ── Slide timing XML lives in ppt/slides/slide?.xml
    # The <p:transition> element controls advance timing.
    spTree = slide.shapes._spTree
    slide_elem = slide._element          # <p:sld>

    # Find or create <p:transition>
    trans = slide_elem.find(qn("p:transition"))
    if trans is None:
        trans = etree.SubElement(slide_elem, qn("p:transition"))

    # advTm = advance time in milliseconds; advClick="0" disables click advance
    trans.set("advTm",    str(SLIDE_DURATION_MS))
    trans.set("advClick", "0")
    trans.set("spd",      "fast")

    print(f"  Slide {i+1:02d} — timing set to {SLIDE_DURATION_MS} ms")

prs.save(PPTX_PATH)
print(f"\nSaved: {PPTX_PATH}")
print(f"Total duration: {len(prs.slides) * SLIDE_DURATION_MS / 1000:.1f} seconds")
